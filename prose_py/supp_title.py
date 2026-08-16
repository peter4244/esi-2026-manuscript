"""Supplement title + auto-generated TOC.

Takes toc_entries collected by build_supplement.collect_toc_entries and emits
the "Supplementary Material" title, subtitle, and body-text TOC linked to
each supplement table's legend bookmark.
"""
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

from tables.docx_helpers import FONT_NAME, BODY_FS, body_toc


TITLE = "Supplementary Material"
SUBTITLE = ("The Emphysema Severity Index: A Spirometric Representation of "
            "CT-Defined Structural Abnormalities in a Multidimensional COPD "
            "Framework")


def build(doc, toc_entries=None):
    # Title
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.name = FONT_NAME
    r.font.size = Pt(BODY_FS + 4)

    # Subtitle
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run(SUBTITLE)
    r.italic = True
    r.font.name = FONT_NAME
    r.font.size = Pt(BODY_FS)

    doc.add_paragraph("")

    if toc_entries:
        body_toc(doc, toc_entries)
