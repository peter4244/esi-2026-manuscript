"""Top-level orchestrator — v9 supplement build.

Two-pass:
  Pass 1: import every supp module (no side effects at import required by
          contract) and collect (TABLE_NUM, TITLE) → TOC entries.
  Pass 2: walk SUPP_PROSE_ORDER (title/TOC + optional frontmatter) then call
          each supp module's build(doc) in manifest order.

Per-artifact try/except; failures render as red paragraphs. TOC-pass failures
render as [TOC ENTRY UNAVAILABLE].

Usage:
    python build_supplement.py
"""
import sys, os, importlib, traceback
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from manifest import (SUPP_PROSE_ORDER, SUPP_TABLES, SUPPLEMENT_OUT, BUILD_LOG_SUPP)
from tables.docx_helpers import (
    FONT_NAME, BODY_FS, resolve_markers, render_markdown_to_docx,
    body_toc, format_label,
)


def init_document():
    doc = Document()
    for section in doc.sections:
        section.page_width  = Inches(8.5); section.page_height = Inches(11)
        section.top_margin  = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin  = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]; normal.font.name = FONT_NAME; normal.font.size = Pt(BODY_FS)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"): rFonts.set(qn(k), FONT_NAME)
    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        st = styles[name]
        st.font.name = FONT_NAME; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.0
    return doc


class BuildLog:
    def __init__(self, path):
        self.path = path; self.entries = []
    def note(self, kind, name, msg):
        self.entries.append(f"[{kind}] {name}: {msg}")
    def save(self):
        with open(self.path, "a") as f:
            f.write("--- supplement build ---\n")
            f.write("\n".join(self.entries) + "\n")


def red_paragraph(doc, msg):
    p = doc.add_paragraph()
    r = p.add_run(f"[{msg}]"); r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.name = FONT_NAME; r.font.size = Pt(BODY_FS)
    return p


# ---------------------------------------------------------------------------
# Pass 1: TOC generation
# ---------------------------------------------------------------------------
def collect_toc_entries(log):
    """Import each supp module without executing build(). Read TABLE_NUM and
    TITLE. Return a list of {num, anchor, title} for body_toc()."""
    entries = []
    for name in SUPP_TABLES:
        try:
            mod = importlib.import_module(f"tables.supp.{name}")
            num   = getattr(mod, "TABLE_NUM")
            title = getattr(mod, "TITLE")
            anchor = f"tbl_{num.replace('.', '_')}"
            entries.append({"num": num, "anchor": anchor, "title": title})
            log.note("OK", f"toc:{name}", f"TABLE_NUM={num}, TITLE={title}")
        except Exception as e:
            entries.append({
                "num": "??", "anchor": "broken",
                "title": f"[TOC ENTRY UNAVAILABLE: {name} — {e}]"
            })
            log.note("FAIL", f"toc:{name}", str(e))
    return entries


# ---------------------------------------------------------------------------
# Pass 2: render tables + prose
# ---------------------------------------------------------------------------
def build_prose_section(doc, section_label, section_source, insertion_registry,
                        insertion_callbacks, toc_entries, log):
    try:
        if section_source.startswith("prose_py."):
            mod = importlib.import_module(section_source)
            # supp title/toc modules take toc_entries as a kwarg
            try:
                mod.build(doc, toc_entries=toc_entries)
            except TypeError:
                mod.build(doc)
        elif section_source.endswith(".md"):
            src_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                     section_source)
            if not os.path.exists(src_path):
                log.note("SKIP", section_label, f"file missing: {src_path}")
                return
            with open(src_path) as f:
                md_text = f.read()
            if not md_text.strip():
                log.note("SKIP", section_label, "empty file")
                return
            resolved, _ = resolve_markers(md_text, insertion_registry)
            render_markdown_to_docx(doc, resolved,
                                    insertion_callbacks=insertion_callbacks)
        else:
            raise ValueError(f"Unknown prose source: {section_source}")
        log.note("OK", section_label, "rendered")
    except Exception as e:
        red_paragraph(doc, f"BUILD FAILED: {section_label} — {e}")
        log.note("FAIL", section_label, str(e))


def build_table(doc, table_name, log):
    try:
        mod = importlib.import_module(f"tables.supp.{table_name}")
        mod.build(doc)
        log.note("OK", table_name, "table rendered")
    except Exception as e:
        red_paragraph(doc, f"BUILD FAILED: {table_name} — {e}")
        log.note("FAIL", table_name, str(e))


def main():
    log = BuildLog(BUILD_LOG_SUPP)
    doc = init_document()

    # Pass 1: TOC
    print("Pass 1 — collecting TOC entries")
    toc_entries = collect_toc_entries(log)

    # Prose sections (title + auto-TOC) — the supp title module calls body_toc
    # if it wants; otherwise we emit it here after the title block.
    #
    # Supplement prose currently has NO {{INSERT:...}} markers by design:
    # supplement tables are rendered by the pass-2 iteration below, not by
    # inline prose markers. If a future supplement prose section wants to embed
    # a supplement table via {{INSERT:...}}, populate this registry.
    insertion_registry = set()
    insertion_callbacks = {}
    for label, source in SUPP_PROSE_ORDER:
        print(f"[prose] {label} ← {source}")
        build_prose_section(doc, label, source, insertion_registry,
                            insertion_callbacks, toc_entries, log)

    # If no prose_py.supp_title exists yet, emit the TOC directly.
    if not any(src.startswith("prose_py.") for _, src in SUPP_PROSE_ORDER):
        body_toc(doc, toc_entries)

    # Pass 2: render each table in manifest order
    print("Pass 2 — rendering tables")
    for name in SUPP_TABLES:
        print(f"[table] {name}")
        build_table(doc, name, log)

    doc.save(SUPPLEMENT_OUT)
    log.save()
    print(f"\nSaved: {SUPPLEMENT_OUT}")
    n_ok   = sum(1 for e in log.entries if e.startswith("[OK]"))
    n_fail = sum(1 for e in log.entries if e.startswith("[FAIL]"))
    print(f"Build summary: {n_ok} OK, {n_fail} FAIL")


if __name__ == "__main__":
    main()
