#!/usr/bin/env python3
"""Build the CT-free manuscript .docx from source.

Unlike the v15 build, this one stays runnable. Its prose lives in METHODS.md
and RESULTS.md, its tables are constructed from the artifacts in assets/ rather
than typed, and its figures are the PNGs the figure scripts emit. Nothing here
is hand-edited into the .docx, so rebuilding never destroys anything, which is
what retired build_manuscript.py in the parent directory.

Prose files carry two kinds of editorial marking that must not reach the
document: claim ids in braces, {RISK-01}, which tie a sentence to its registry
entry, and provenance notes, **[v15 para 67, verbatim]**, which record reuse.
Both are stripped here so the sources stay annotated while the output is clean.

Usage, from anywhere:
    python3 /abs/path/to/ctfree/build_manuscript.py
"""
import csv
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from docx import Document                                    # noqa: E402
from docx.shared import Pt, Inches, RGBColor                 # noqa: E402
from docx.oxml import OxmlElement                            # noqa: E402
from docx.oxml.ns import qn                                  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                # noqa: E402
from docx.enum.table import WD_ALIGN_VERTICAL                # noqa: E402

from tables.docx_helpers import FONT_NAME, CONTENT_WIDTH_IN  # noqa: E402
import pjc_guard
import prose_guard                                             # noqa: E402
import display_order                                           # noqa: E402
from supp_registry import supp_num                             # noqa: E402

ASSETS = os.path.join(HERE, "assets")
FIGS   = os.path.join(HERE, "figures")
OUT    = os.path.join(HERE, "manuscript",
                      "CT-free MD-COPD manuscript draft v1.docx")
BODY_FS, TBL_FS = 11, 9
CATS = ["noCOPD", "AFL-only", "COPD-minor", "COPD-major"]
# Mortality and exacerbation models are fitted on different participant sets,
# so the denominator shown must follow the outcome rather than being one number
# reused for all three.
N_COL = {"all": "n_mort", "resp": "n_mort", "exac": "n_exac"}
# Short names agreed for the manuscript, defined in Methods and restated in
# Results. They replace the numbered "schema N" labels, which required the
# reader to hold a lookup table to parse a sentence.
SCHEMA_NAME = {"S1": "Fixed ratio", "S2": "MD-COPD",
               "S3": "NoCT classification", "S4": "ESI classification"}

CLAIM_ID   = re.compile(r"\s*\{[A-Z][A-Za-z0-9-]*(?:,\s*[A-Z][A-Za-z0-9-]*)*\}")
# Any bracketed bold line is an editorial provenance note, not prose.
PROVENANCE = re.compile(r"^\*\(.*\)\*$|^\*\*\[.*\]\*\*$")


def load(name):
    with open(os.path.join(ASSETS, name)) as f:
        return list(csv.DictReader(f))


def init_document():
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_FS)
    rPr = normal.element.get_or_add_rPr()
    rF = rPr.find(qn("w:rFonts"))
    if rF is None:
        rF = OxmlElement("w:rFonts"); rPr.append(rF)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rF.set(qn(k), FONT_NAME)
    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT_NAME, Pt(size), True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.0
    return doc


def emphasis(par, text, size=BODY_FS):
    """Render **bold** spans; everything else plain."""
    for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not chunk:
            continue
        r = par.add_run(chunk[2:-2] if chunk.startswith("**") else chunk)
        r.bold = chunk.startswith("**")
        r.font.name, r.font.size = FONT_NAME, Pt(size)


def add_prose(doc, md_path, skip_after=None):
    """Render a prose source, stopping at `skip_after` so the editorial
    'Still to write' and 'Open' sections never reach the document."""
    with open(md_path) as f:
        lines = f.read().split("\n")
    # Everything before the first section heading is editorial front matter
    # explaining the annotations, and belongs in the source rather than the
    # document.
    first = next((i for i, l in enumerate(lines) if l.startswith("## ")), 0)
    lines = lines[first:]
    para, emitted = [], 0

    def flush():
        nonlocal para, emitted
        if not para:
            return
        text = CLAIM_ID.sub("", " ".join(para)).strip()
        para = []
        if not text or PROVENANCE.match(text):
            return
        emphasis(doc.add_paragraph(), text)
        emitted += 1

    for ln in lines:
        t = ln.rstrip()
        if t.startswith("#"):
            title = t.lstrip("#").strip()
            if skip_after and title.lower().startswith(skip_after.lower()):
                flush()
                return emitted
            flush()
            if t.startswith("## ") and title.lower() not in ("introduction", "discussion"):
                doc.add_paragraph(title, style="Heading 2")
        elif t.startswith("---") or t.startswith("|") or t.startswith("```"):
            flush()
        elif not t.strip():
            flush()
        else:
            para.append(t.strip())
    flush()
    return emitted


def set_cell_margins(tbl, inches):
    """Word's default cell padding is 0.08 inch a side, which is 0.16 inch of
    every column lost to whitespace. At nine columns that is the difference
    between a number fitting on one line and wrapping mid-value."""
    tblPr = tbl._element.find(qn("w:tblPr"))
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar"); tblPr.append(mar)
    for side in ("top", "left", "bottom", "right"):
        el = mar.find(qn("w:" + side))
        if el is None:
            el = OxmlElement("w:" + side); mar.append(el)
        el.set(qn("w:w"), str(int(round(inches * 1440))))
        el.set(qn("w:type"), "dxa")


class Section(str):
    """A table row that names the group of rows beneath it, spanning every column."""


def add_table(doc, headers, rows, widths, body_fs=TBL_FS, pad_in=0.05):
    assert abs(sum(widths) - CONTENT_WIDTH_IN) < 0.02, sum(widths)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    grid = t._element.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(w * 1440))))
        gc.set(qn("w:type"), "dxa")
    tblPr = t._element.find(qn("w:tblPr"))
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    set_cell_margins(t, pad_in)
    for i, row in enumerate([headers] + rows):
        if isinstance(row, Section):
            cell = t.rows[i].cells[0].merge(t.rows[i].cells[-1])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(row))
            r.bold = True
            r.font.name, r.font.size = FONT_NAME, Pt(body_fs)
            continue
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.width = Inches(widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if j:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.bold = (i == 0)
            r.font.name, r.font.size = FONT_NAME, Pt(body_fs)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return t


def legend(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    r.font.name, r.font.size = FONT_NAME, Pt(BODY_FS - 1)
    r2 = p.add_run(text)
    r2.font.name, r2.font.size = FONT_NAME, Pt(BODY_FS - 1)


# --------------------------------------------------------------------------
# Tables, built from the artifacts
# --------------------------------------------------------------------------
def t_classifications(doc, label="Table 1."):
    lab = {(r["schema"], r["category"]): int(r["n"]) for r in load("schema_labels.csv")}
    n = lambda s, c: f"{lab[(s, c)]:,}" if (s, c) in lab else "—"
    n_total = sum(lab[(s, c)] for (s, c) in lab if s == "S2")
    # ESI threshold read from the fitted-schema artifact so this cannot drift
    # from the actual model.
    fit = {r["schema"]: r for r in load("schema_fit.csv")}
    esi_t = float(fit["S4"]["t_low"])
    esi_desc = f"ESI ≥ {esi_t:.2f}, dyspnea, SGRQ, chronic bronchitis (≥2)"
    rows = [
        ["Fixed ratio", "none; COPD is airflow limitation",
         n("S1", "noCOPD"), "—", "—", n("S1", "COPD")],
        ["MD-COPD", "emphysema, wall thickening, dyspnea, SGRQ, chronic bronchitis (≥3)",
         n("S2", "noCOPD"), n("S2", "AFL-only"), n("S2", "COPD-minor"), n("S2", "COPD-major")],
        ["NoCT classification", "dyspnea, SGRQ, chronic bronchitis (≥2)",
         n("S3", "noCOPD"), n("S3", "AFL-only"), n("S3", "COPD-minor"), n("S3", "COPD-major")],
        ["ESI classification", esi_desc,
         n("S4", "noCOPD"), n("S4", "AFL-only"), n("S4", "COPD-minor"), n("S4", "COPD-major")]]
    add_table(doc, ["Classification", "Minor criteria", "noCOPD", "AFL-only",
                    "COPD-minor", "COPD-major"],
              rows, [1.25, 1.95, 0.72, 0.92, 0.84, 0.82])
    legend(doc, label,
           f"The four classifications applied to the same {n_total:,} participants.")


def t_risk_common_ref(doc, label="Table 3."):
    """Adjusted risk of every group of every multidimensional classification
    against the common reference: the participants all three assign to noCOPD.
    Adjusted only (Pete, 2026-09-13); the crude ratios for the same groups are
    in the group comparison table. A cell under the event floor keeps its point
    estimate and loses its interval."""
    adj = {(r["schema"], r["category"]): r for r in load("consensus_ref_risk.csv")}
    ref = load("consensus_ref_group.csv")[0]
    NM_TITLE = {"S2": "MD-COPD", "S3": "NoCT classification", "S4": "ESI classification"}
    ORDER = ["AFL-only", "COPD-minor", "COPD-major"]
    FLAG, TRUE = "\u2020", ("TRUE", "T")

    def cell(r, est, lo, hi, few):
        if few:
            return f"{float(r[est]):.2f}{FLAG}"
        return f"{float(r[est]):.2f} ({float(r[lo]):.2f}\u2013{float(r[hi]):.2f})"

    rows = [["Common reference (noCOPD under all three)", f"{int(ref['deaths']):,}", "reference",
             f"{int(ref['resp_deaths']):,}", "reference", "reference"]]
    for s_ in ("S2", "S3", "S4"):
        rows.append(Section(NM_TITLE[s_]))
        for g in ORDER:
            r = adj[(s_, g)]
            rows.append([
                "\u2003" + g, f"{int(r['deaths']):,}",
                cell(r, "all_HR", "all_LCI", "all_UCI", r["all_few_events"].upper() in TRUE),
                f"{int(r['resp_deaths']):,}",
                cell(r, "resp_HR", "resp_LCI", "resp_UCI", r["resp_few_events"].upper() in TRUE),
                cell(r, "exac_IRR", "exac_LCI", "exac_UCI", False)])
    add_table(doc, ["Group", "Deaths", "All-cause HR (95% CI)", "Resp. deaths",
                    "Resp. HR (95% CI)", "Exacerbation IRR (95% CI)"],
              rows, [2.05, 0.65, 1.25, 0.65, 1.00, 0.90])
    flagged = any(FLAG in str(v) for r_ in rows if not isinstance(r_, Section) for v in r_)
    legend(doc, label,
           "Adjusted risk of each group of the three multidimensional classifications against "
           f"a common reference, the {int(ref['n_cohort']):,} participants all three "
           "classifications assign to noCOPD. Adjusted models carry age, sex, race, current "
           "smoking status, pack-years and body mass index, with prior exacerbation frequency "
           "added for exacerbations. Crude rate ratios for the same groups are given in "
           f"Supplemental Table {supp_num('groups_vs_mdcopd')}. "
           + (f"{FLAG} fewer than 10 events in the group: the estimate is given without an "
              "interval. " if flagged else "")
           + "AFL-only, airflow limitation without other criteria; HR, hazard ratio; IRR, "
           "incidence rate ratio; CI, confidence interval.")

def t_discord_risk(doc, label="Table 4."):
    """Cross-classification of MD-COPD and the ESI classification: each group
    against its own stratum's reference, crude ratios only, respiratory mortality
    out for too few deaths (Pete, 2026-09-11),
    as in Figure 6. One table with a header row per stratum; each stratum
    carries its own reference row, because the references differ (both call
    noCOPD where spirometry is preserved, both call AFL-only where it is not).
    The event floor applies to the group and to its reference alike."""
    FLOOR, FLAG = 10, "\u2020"
    STRATA = [("Preserved spirometry", "discord_rates.csv", "Both-noCOPD"),
              ("Airflow limitation", "discord_rates_afl.csv", "Both-AFL-only")]
    ORDER = ["CT-only-COPD", "ESI-only-COPD", "Both-COPD"]
    cru = {(r["stratum"], r["group"], r["outcome"]): r for r in load("discord_crude_strata.csv")}

    def cell(r, n_ev, n_ref):
        if n_ev == 0:
            return "not estimable"
        if n_ev < FLOOR or n_ref < FLOOR or r["lo"] in ("", "NA"):
            return f"{float(r['rr']):.2f}{FLAG}"
        return f"{float(r['rr']):.2f} ({float(r['lo']):.2f}\u2013{float(r['hi']):.2f})"

    rows, resp = [], {}
    for title, rf, refg in STRATA:
        rates = {r["group"]: r for r in load(rf)}
        R = rates[refg]
        ref_ev = {"all": int(R["deaths"]), "resp": int(R["resp_deaths"]), "exac": 10 ** 9}
        rows.append(Section(title))
        rows.append([f"\u2003{refg} (reference)", f"{int(R['n']):,}", f"{int(R['deaths']):,}", "1.00",
                     "1.00"])
        resp[title] = [(refg, int(R["resp_deaths"]))]
        for g in ORDER:
            c = {o: cru[(title, g, o)] for o in ("all", "resp", "exac")}
            ev = {o: int(float(c[o]["events"])) for o in c}
            assert ev["all"] == int(rates[g]["deaths"]) and ev["resp"] == int(rates[g]["resp_deaths"]), g
            rows.append(["\u2003" + g, f"{int(rates[g]['n']):,}", f"{ev['all']:,}",
                         cell(c["all"], ev["all"], ref_ev["all"]),
                         cell(c["exac"], ev["exac"], ref_ev["exac"])])
            resp[title].append((g, ev["resp"]))
    # With respiratory mortality out, every cell clears the floor; the legend
    # carries no flag note, so a flagged cell appearing later must stop the build.
    assert not any(FLAG in str(v) for r in rows if not isinstance(r, Section) for v in r), \
        "S7 has a cell under the event floor; restore the flag note in its legend"
    add_table(doc, ["Group", "n", "Deaths", "All-cause mortality RR (95% CI)",
                    "Exacerbation RR (95% CI)"],
              rows, [2.00, 0.70, 0.80, 1.50, 1.50])
    resp_txt = "; ".join(
        f"with {st[0].lower() + st[1:]}, " + ", ".join(f"{n} in {g}" for g, n in v)
        for st, v in resp.items())
    legend(doc, label,
           "Participants cross-classified by MD-COPD and the ESI classification within each "
           "stratum of airflow limitation. CT-only-COPD is COPD under MD-COPD only, "
           "ESI-only-COPD COPD under the ESI classification only, and Both-COPD COPD under "
           "both. Each stratum is estimated against its own reference row, so estimates are "
           "comparable within a stratum and not between strata. Crude rate ratios (RR) are "
           "the observed event rate in the group divided by the rate in the reference; "
           "intervals are exact Poisson intervals for deaths and subject-bootstrap intervals "
           "for exacerbations. "
           "Respiratory mortality is not shown because the numbers of respiratory deaths "
           f"are too small for estimation: {resp_txt}. AFL-only, airflow limitation "
           "without other criteria; CI, confidence interval.")

def t_esi_auc(doc, label="Table 2."):
    """ESI's discrimination of the two visual CT criteria, by stratum. This is
    the claim that explains where an ESI criterion helps. The comparison with
    FEV1/FVC belongs in the Supplement: FEV1/FVC cannot serve as the
    replacement criterion at all, so putting the two side by side in the main
    text invites a comparison the framework does not permit."""
    au = load("esi_ct_auc.csv")
    ORDER = ["All participants", "Airflow limitation", "Preserved spirometry"]
    rows, seen = [], set()
    for crit in ("Visual emphysema", "Airway wall thickening"):
        for st in ORDER:
            r = next(x for x in au if x["criterion"] == crit and x["stratum"] == st)
            rows.append([crit if crit not in seen else "", st,
                         f"{int(r['n']):,}", f"{float(r['prevalence']):.1f}",
                         f"{float(r['auc_ESI']):.2f}"])
            seen.add(crit)
    add_table(doc, ["CT criterion", "Stratum", "n", "Prevalence (%)",
                    "AUC for ESI"],
              rows, [1.40, 1.46, 0.76, 1.10, 1.78])
    legend(doc, label,
           "Discrimination capability of ESI for the two visual CT criteria it "
           "replaces, shown overall and within stratum of airflow limitation. "
           "AUC, area under the receiver operating characteristic curve.")


def figure_width(meta_path, default=CONTENT_WIDTH_IN):
    """Content width a figure was rendered for, from the sidecar its script
    writes. Absent a sidecar the figure is a portrait one."""
    if not os.path.exists(meta_path):
        return default
    for line in open(meta_path):
        k, _, v = line.strip().partition("=")
        if k == "content_width_in":
            w = float(v)
            if not (CONTENT_WIDTH_IN <= w <= 9.0 + 1e-6):
                raise SystemExit(f"{meta_path}: content width {w} outside 6.5 to 9 in")
            return w
    raise SystemExit(f"{meta_path}: no content_width_in")


def new_section(doc, landscape):
    """Start a new page in the given orientation, 1 in margins either way."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = ((Inches(11), Inches(8.5)) if landscape
                                       else (Inches(8.5), Inches(11)))
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(1))
    return sec


def read_legend(stem):
    """Legend text from a figure's sidecar: its title sentence, then its body.
    The title used to be dropped, leaving a legend that opened on panel (A)."""
    lines = [l.strip() for l in open(os.path.join(FIGS, stem + "_legend.md")) if l.strip()]
    m = re.match(r"^\*\*Figure \d+\.\s*(.*?)\*\*$", lines[0])
    if not m:
        raise SystemExit(f"{stem}_legend.md: first line is not a bold figure title")
    return f"{m.group(1).strip()} {CLAIM_ID.sub('', lines[-1])}"


def add_figure(doc, png, label, text, width=CONTENT_WIDTH_IN):
    doc.add_picture(png, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    legend(doc, label, text)


def check_references():
    """The reference list and the citations in the prose must agree, in both
    directions. A dangling citation and an uncited entry are both defects that
    reach reviewers, and renumbering is exactly when they appear."""
    body = ""
    for fn, stop in (("ABSTRACT.md", None), ("INTRODUCTION.md", None),
                     ("METHODS.md", "## Still to write"),
                     ("RESULTS.md", "## Open"), ("DISCUSSION.md", None)):
        t = open(os.path.join(HERE, fn)).read()
        body += t.split(stop)[0] if stop and stop in t else t
    cited = set()
    for m in re.finditer(r"\((\d+(?:\s*[,\u2013-]\s*\d+)*)\)", body):
        for part in m.group(1).split(","):
            part = part.strip()
            if re.fullmatch(r"\d+", part):
                cited.add(int(part))
            else:
                r = re.fullmatch(r"(\d+)\s*[\u2013-]\s*(\d+)", part)
                if r:
                    cited.update(range(int(r.group(1)), int(r.group(2)) + 1))
    listed = set()
    for line in open(os.path.join(HERE, "REFERENCES.md")):
        m = re.match(r"^(\d+)\.\s+\S", line)
        if m:
            listed.add(int(m.group(1)))
    dangling = sorted(cited - listed)
    uncited = sorted(listed - cited)
    gaps = [n for n in range(1, max(listed) + 1) if n not in listed] if listed else []
    if dangling:
        raise SystemExit(f"prose cites {dangling}, absent from REFERENCES.md")
    if gaps:
        raise SystemExit(f"REFERENCES.md skips {gaps}; numbering must be contiguous")
    if uncited:
        print(f"  note: references {uncited} listed but never cited")
    return len(listed)


def read_titlepage():
    """TITLEPAGE.md is the single source of truth for the title, author list,
    affiliations and key words. It is Pete's text; nothing regenerates it."""
    sections, key = {}, None
    for line in open(os.path.join(HERE, "TITLEPAGE.md")):
        line = line.rstrip("\n")
        if line.startswith("<!--") or line.startswith("     ") and key is None:
            continue
        if line.startswith("# "):
            key = line[2:].strip()
            sections[key] = []
        elif key and line.strip():
            sections[key].append(line.strip())
    for k in ("TITLE", "AUTHORS", "AFFILIATIONS", "KEYWORDS"):
        if not sections.get(k):
            raise SystemExit(f"TITLEPAGE.md is missing the {k} section")
    return sections


def add_front_matter(doc):
    tp = read_titlepage()

    p = doc.add_paragraph()
    r = p.add_run(" ".join(tp["TITLE"]))
    r.bold = True
    r.font.name, r.font.size = FONT_NAME, Pt(12)

    p = doc.add_paragraph()
    r = p.add_run(" ".join(tp["AUTHORS"]))
    r.font.name, r.font.size = FONT_NAME, Pt(BODY_FS)

    doc.add_paragraph()
    for aff in tp["AFFILIATIONS"]:
        p = doc.add_paragraph()
        r = p.add_run(aff)
        r.font.name, r.font.size = FONT_NAME, Pt(BODY_FS)

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Key words:")
    r.bold = True
    r.font.name, r.font.size = FONT_NAME, Pt(BODY_FS)
    r = p.add_run(" " + " ".join(tp["KEYWORDS"]))
    r.font.name, r.font.size = FONT_NAME, Pt(BODY_FS)


def heading_on_new_page(doc, text):
    """Start a section on its own page. page_break_before travels with the
    heading, so reflow above it cannot leave the break stranded mid-page the
    way a separately inserted break can."""
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = True
    return p


def main():
    pjc_guard.check("The manuscript build")
    prose_guard.check("The manuscript build")
    prose_guard.check_owned("The manuscript build")
    display_order.check(HERE, {"Table": ["1", "2"],
                               "Figure": [str(i) for i in range(1, 7)]},
                        "The manuscript build")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc = init_document()

    add_front_matter(doc)

    heading_on_new_page(doc, "ABSTRACT")
    add_prose(doc, os.path.join(HERE, "ABSTRACT.md"))
    heading_on_new_page(doc, "INTRODUCTION")
    n_i = add_prose(doc, os.path.join(HERE, "INTRODUCTION.md"))
    doc.add_paragraph("METHODS", style="Heading 1")
    n_m = add_prose(doc, os.path.join(HERE, "METHODS.md"), skip_after="Still to write")
    doc.add_paragraph("RESULTS", style="Heading 1")
    n_r = add_prose(doc, os.path.join(HERE, "RESULTS.md"), skip_after="Open")

    doc.add_paragraph("DISCUSSION", style="Heading 1")
    n_d = add_prose(doc, os.path.join(HERE, "DISCUSSION.md"))

    doc.add_paragraph("REFERENCES", style="Heading 1")
    n_ref = check_references()
    for line in open(os.path.join(HERE, "REFERENCES.md")):
        if re.match(r"^\d+\.\s+\S", line):
            doc.add_paragraph(line.strip())

    doc.add_page_break()
    doc.add_paragraph("TABLES", style="Heading 1")
    t_classifications(doc, "Table 1.")
    doc.add_paragraph()
    t_esi_auc(doc, "Table 2.")

    # Figure 1 declares its own content width. A figure rendered for a landscape
    # page must be placed at that width, or every font in it lands smaller than
    # the size the figure was validated at.
    fig1_w = figure_width(os.path.join(FIGS, "figure1_flow.meta"))
    landscape = fig1_w > CONTENT_WIDTH_IN + 0.01
    if landscape:
        new_section(doc, landscape=True)
    else:
        doc.add_page_break()
    doc.add_paragraph("FIGURES", style="Heading 1")
    add_figure(doc, os.path.join(FIGS, "figure1_flow.png"), "Figure 1.",
               read_legend("figure1_flow"), width=fig1_w)
    if landscape:
        new_section(doc, landscape=False)
    ref = load("consensus_ref_group.csv")[0]
    for i, (png, grp, note) in enumerate([
            ("figure2_aflonly.png", "AFL-only",
             "This is the group that withholds a COPD diagnosis from "
             "participants with airflow limitation, so it is the group a "
             "CT-free classification has to keep at low risk."),
            ("figure3_copdminor.png", "COPD-minor",
             "This is the group a replacement criterion must reach among "
             "participants with preserved spirometry."),
            ("figure4_copdmajor.png", "COPD-major",
             "Estimates are largest under the NoCT classification because that "
             "classification moved 833 of this group's members into AFL-only, "
             "leaving a smaller and more severe group behind.")], start=2):
        title = ("Prognostic risks for AFL-only pathway subjects in each "
                 "classification. " if i == 2 else "")          # Pete, 2026-09-11
        add_figure(doc, os.path.join(FIGS, png), f"Figure {i}.",
                   f"{title}Crude (open circles) and adjusted (filled circles) risk "
                   f"for the {grp} group under each multidimensional "
                   f"classification, against a reference common to all three: "
                   f"the {int(ref['n_cohort']):,} participants every "
                   f"classification assigns to noCOPD. A common reference makes "
                   f"an estimate under one classification comparable with an "
                   f"estimate under another. Panels carry separate x scales "
                   f"because the three outcomes differ in magnitude; the "
                   f"comparison the figure supports is between classifications "
                   f"within a panel. A point drawn without an interval had "
                   f"fewer than 10 events in that group, so no interval was "
                   f"estimated. {note} AFL-only, airflow limitation without "
                   f"other criteria.")

    # Figures 5 and 6 carry their legends in a sidecar, as Figure 1 does, and
    # declare their own widths; the page turns where a width needs it to.
    landscape_now = False
    for n, stem in ((5, "figure5_profile"), (6, "figure6_discord_risk")):
        w = figure_width(os.path.join(FIGS, stem + ".meta"))
        want = w > CONTENT_WIDTH_IN + 0.01
        if want != landscape_now:
            new_section(doc, landscape=want)
            landscape_now = want
        add_figure(doc, os.path.join(FIGS, stem + ".png"), f"Figure {n}.",
                   read_legend(stem), width=w)
    if landscape_now:
        new_section(doc, landscape=False)

    doc.save(OUT)
    print(f"wrote {OUT}\n  {n_i} Introduction, {n_m} Methods, {n_r} Results, "
          f"{n_d} Discussion paragraphs, {n_ref} references")
    # A floor against a source file failing to render, not a target length.
    # Pete's Introduction is four paragraphs by choice.
    assert n_i >= 4 and n_m > 10 and n_r > 10 and n_d > 5, \
        "prose came out suspiciously short"


if __name__ == "__main__":
    main()
